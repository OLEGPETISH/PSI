from __future__ import annotations

import random
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook


JOURNAL_PATH = Path(__file__).with_name("journal.xlsx")
HEADERS = ("Last name", "First name", "Номер билета", "Дата и время")


def read_line(prompt: str) -> str | None:
    """Read a console line and return None if ESC is pressed."""
    if sys.platform != "win32":
        raise RuntimeError("Перехват ESC поддерживается в Windows.")

    import msvcrt

    sys.stdout.write(prompt)
    sys.stdout.flush()
    characters: list[str] = []

    while True:
        char = msvcrt.getwch()
        if char == "\x1b":
            sys.stdout.write("\n")
            return None
        if char in ("\r", "\n"):
            sys.stdout.write("\n")
            return "".join(characters)
        if char in ("\x00", "\xe0"):
            msvcrt.getwch()
            continue
        if char == "\b":
            if characters:
                characters.pop()
                sys.stdout.write("\b \b")
                sys.stdout.flush()
            continue
        if char.isprintable():
            characters.append(char)
            sys.stdout.write(char)
            sys.stdout.flush()


def ask_non_empty(prompt: str) -> str | None:
    while True:
        value = read_line(prompt)
        if value is None:
            return None
        value = value.strip()
        if value:
            return value
        print("Поле не может быть пустым. Попробуйте ещё раз.")


def append_student(
    path: Path,
    last_name: str,
    first_name: str,
    ticket_number: int,
    recorded_at: datetime | None = None,
) -> None:
    recorded_at = recorded_at or datetime.now()
    if path.exists():
        workbook = load_workbook(path)
        worksheet = workbook.active
    else:
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Журнал"
        worksheet.append(HEADERS)

    try:
        worksheet.append(
            (
                last_name,
                first_name,
                ticket_number,
                recorded_at.strftime("%Y-%m-%d %H:%M:%S"),
            )
        )
        workbook.save(path)
    finally:
        workbook.close()


def save_student(last_name: str, first_name: str, ticket_number: int) -> bool:
    while True:
        try:
            append_student(JOURNAL_PATH, last_name, first_name, ticket_number)
            return True
        except (PermissionError, OSError) as error:
            print(
                f"\nНе удалось сохранить запись в {JOURNAL_PATH.name}: {error}\n"
                "Закройте файл в Excel и нажмите Enter для повторной попытки "
                "(ESC — отменить запись)."
            )
            if read_line("> ") is None:
                print("Запись не сохранена.")
                return False


def main() -> None:
    print("Генератор экзаменационных билетов")
    print("Для выхода нажмите ESC.\n")

    while True:
        last_name = ask_non_empty("Last name: ")
        if last_name is None:
            print("Работа приложения завершена.")
            return

        first_name = ask_non_empty("First name: ")
        if first_name is None:
            print("Работа приложения завершена.")
            return

        ticket_number = random.randint(1, 20)
        print(f"Билет № {ticket_number}")
        if save_student(last_name, first_name, ticket_number):
            print(f"Запись сохранена в {JOURNAL_PATH.name}.\n")


if __name__ == "__main__":
    try:
        main()
    except KeyboardInterrupt:
        print("\nРабота приложения прервана.")

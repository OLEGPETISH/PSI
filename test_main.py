import tempfile
import unittest
from contextlib import redirect_stdout
from datetime import datetime
from io import StringIO
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

from main import HEADERS, append_student, ask_non_empty, save_student


class AppendStudentTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory()
        self.journal = Path(self.temp_dir.name) / "journal.xlsx"
        self.timestamp = datetime(2026, 10, 21, 12, 0, 0)

    def tearDown(self) -> None:
        self.temp_dir.cleanup()

    def test_creates_workbook_with_headers_and_student(self) -> None:
        append_student(self.journal, "Иванов", "Иван", 7, self.timestamp)

        workbook = load_workbook(self.journal, read_only=True)
        rows = list(workbook.active.values)
        workbook.close()

        self.assertEqual(rows[0], HEADERS)
        self.assertEqual(
            rows[1],
            ("Иванов", "Иван", 7, "2026-10-21 12:00:00"),
        )

    def test_appends_without_changing_existing_rows(self) -> None:
        append_student(self.journal, "Иванов", "Иван", 7, self.timestamp)
        append_student(self.journal, "Петрова", "Анна", 12, self.timestamp)

        workbook = load_workbook(self.journal, read_only=True)
        rows = list(workbook.active.values)
        workbook.close()

        self.assertEqual(len(rows), 3)
        self.assertEqual(rows[1][0:3], ("Иванов", "Иван", 7))
        self.assertEqual(rows[2][0:3], ("Петрова", "Анна", 12))

    @patch("main.read_line", side_effect=["   ", "  Иванов  "])
    def test_empty_name_is_rejected_and_whitespace_is_trimmed(self, read_line) -> None:
        with redirect_stdout(StringIO()):
            name = ask_non_empty("Last name: ")

        self.assertEqual(name, "Иванов")
        self.assertEqual(read_line.call_count, 2)

    @patch("main.read_line", return_value="")
    @patch("main.append_student", side_effect=[PermissionError("file is locked"), None])
    def test_locked_journal_is_retried(self, append, read_line) -> None:
        with patch("main.JOURNAL_PATH", self.journal), redirect_stdout(StringIO()):
            saved = save_student("Иванов", "Иван", 7)

        self.assertTrue(saved)
        self.assertEqual(append.call_count, 2)
        read_line.assert_called_once()


if __name__ == "__main__":
    unittest.main()
